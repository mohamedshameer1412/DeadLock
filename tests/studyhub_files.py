"""Builders for the sample files StudyHub's tests upload. Everything is made in memory; nothing is downloaded."""
from __future__ import annotations

import io


def _esc(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(pages: list[str], outline: list[tuple[str, int]] | None = None, title: str | None = None) -> bytes:
    """A small but valid text PDF. `pages` are texts (\\n = new line); `outline` = [(title, 0-based page index)]."""
    objs: dict[int, bytes] = {}
    n_pages = len(pages)
    first_page_obj = 4
    kids = " ".join(f"{first_page_obj + 2 * i} 0 R" for i in range(n_pages))
    outline_root = first_page_obj + 2 * n_pages
    catalog = "<< /Type /Catalog /Pages 2 0 R" + (f" /Outlines {outline_root} 0 R" if outline else "") + " >>"
    objs[1] = catalog.encode()
    objs[2] = f"<< /Type /Pages /Kids [{kids}] /Count {n_pages} >>".encode()
    objs[3] = b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    for i, text in enumerate(pages):
        page_obj, content_obj = first_page_obj + 2 * i, first_page_obj + 2 * i + 1
        objs[page_obj] = (f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_obj} 0 R "
                          "/Resources << /Font << /F1 3 0 R >> >> >>").encode()
        lines = ["BT", "/F1 11 Tf", "14 TL", "50 740 Td"]
        for line in text.split("\n"):
            lines.append(f"({_esc(line)}) Tj T*" if line else "T*")
        lines.append("ET")
        stream = "\n".join(lines).encode("latin-1", "replace")
        objs[content_obj] = b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream"
    if outline:
        item0 = outline_root + 1
        objs[outline_root] = f"<< /Type /Outlines /First {item0} 0 R /Last {item0 + len(outline) - 1} 0 R /Count {len(outline)} >>".encode()
        for j, (name, page_idx) in enumerate(outline):
            parts = [f"/Title ({_esc(name)})", f"/Parent {outline_root} 0 R",
                     f"/Dest [{first_page_obj + 2 * page_idx} 0 R /Fit]"]
            if j > 0:
                parts.append(f"/Prev {item0 + j - 1} 0 R")
            if j < len(outline) - 1:
                parts.append(f"/Next {item0 + j + 1} 0 R")
            objs[item0 + j] = ("<< " + " ".join(parts) + " >>").encode()
    if title:
        info_obj = max(objs) + 1
        objs[info_obj] = f"<< /Title ({_esc(title)}) >>".encode()
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = {}
    for num in sorted(objs):
        offsets[num] = out.tell()
        out.write(b"%d 0 obj\n" % num + objs[num] + b"\nendobj\n")
    xref = out.tell()
    size = max(objs) + 1
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % size)
    for num in range(1, size):
        out.write(b"%010d 00000 n \n" % offsets[num])
    trailer = f"trailer\n<< /Size {size} /Root 1 0 R" + (f" /Info {info_obj} 0 R" if title else "") + " >>\n"
    out.write(trailer.encode() + b"startxref\n%d\n%%%%EOF\n" % xref)
    return out.getvalue()


def make_blank_pdf(pages: int = 2) -> bytes:
    """A PDF whose pages carry no text at all - what a scan looks like to a text extractor."""
    from pypdf import PdfWriter
    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(612, 792)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def encrypt_pdf(data: bytes, password: str = "secret") -> bytes:
    from pypdf import PdfReader, PdfWriter
    w = PdfWriter(clone_from=PdfReader(io.BytesIO(data)))
    w.encrypt(password)
    buf = io.BytesIO()
    w.write(buf)
    return buf.getvalue()


def make_docx(items: list[tuple[str, str]], table: list[list[str]] | None = None, title: str | None = None) -> bytes:
    """items = [(style, text)] with style 'Heading 1'.. or 'Normal'. An optional table is appended at the end."""
    from docx import Document
    d = Document()
    if title:
        d.core_properties.title = title
    for style, text in items:
        if style.startswith("Heading"):
            d.add_heading(text, level=int(style.split()[1]))
        else:
            d.add_paragraph(text)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


SAMPLE_TXT = """# Data Structures

## Stacks

A stack is a last-in first-out collection. The push operation adds an element to the top of the stack and the pop
operation removes the element from the top. Stacks are used to implement function calls and undo features.

## Queues

A queue is a first-in first-out collection. The enqueue operation adds an element at the rear and the dequeue
operation removes the element at the front. Queues are used for scheduling tasks and for breadth-first search.

## Trees

A binary tree is a hierarchical structure in which every node has at most two children, called the left child and
the right child. Inorder traversal visits the left subtree, then the node, then the right subtree.
"""
